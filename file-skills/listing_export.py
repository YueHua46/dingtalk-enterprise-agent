"""Deterministic Listing renderer in the isolated file worker. No model or network calls.

The production worker has python-docx/openpyxl and no artifact-tool installation.
Every spreadsheet cell is explicitly a string; external text cannot become a formula.
"""
import json
import sys
import hashlib
import os
import tempfile
import zipfile
import re
import unicodedata
from difflib import SequenceMatcher
from datetime import datetime
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

RENDER_VERSION = 'listing-report-design-v2'
NAVY, BLUE, MUTED, LIGHT, AMBER = '183153', '235A91', '667085', 'F3F6FA', 'FFF4DC'


def shade(cell, color):
    element = OxmlElement('w:shd')
    element.set(qn('w:fill'), color)
    cell._tc.get_or_add_tcPr().append(element)


def table_style(table, widths=None):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    if widths:
        for column, width in zip(table.columns, widths):
            column.width = Inches(width)
    for row in table.rows:
        no_split = OxmlElement('w:cantSplit')
        row._tr.get_or_add_trPr().append(no_split)
        for i, cell in enumerate(row.cells):
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            if widths:
                cell.width = Inches(widths[i])
            margins = OxmlElement('w:tcMar')
            for edge in ['top', 'bottom', 'left', 'right']:
                el = OxmlElement('w:' + edge)
                el.set(qn('w:w'), '90')
                el.set(qn('w:type'), 'dxa')
                margins.append(el)
            cell._tc.get_or_add_tcPr().append(margins)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(3)
                for run in paragraph.runs:
                    run.font.size = Pt(9)
    # Header repeats if a long source/check table spans pages.
    header = OxmlElement('w:tblHeader')
    table.rows[0]._tr.get_or_add_trPr().append(header)


def source_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    for cell, text in zip(table.rows[0].cells, headers):
        cell.text = text
        shade(cell, NAVY)
        for run in cell.paragraphs[0].runs:
            run.font.color.rgb = RGBColor.from_string('FFFFFF')
            run.bold = True
    for index, values in enumerate(rows):
        cells = table.add_row().cells
        for cell, text in zip(cells, values):
            cell.text = text or '无'
            if index % 2 == 0:
                shade(cell, LIGHT)
    table_style(table, widths)
    return table


def keep_small_table(table, rows):
    # Keep short reference tables together; long inputs retain normal pagination.
    if len(rows) <= 18 and sum(len(value) for row in rows for value in row) <= 4000:
        for row in table.rows[:-1]:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.keep_with_next = True


def changed_text(paragraph, before, after):
    # Bounded literal diff, never infer improvements or alter the copied text.
    if max(len(before), len(after)) > 5000:
        paragraph.add_run(after)
        return
    for kind, _, _, start, end in SequenceMatcher(None, before, after, autojunk=False).get_opcodes():
        run = paragraph.add_run(after[start:end])
        if kind in ['insert', 'replace']:
            run.bold = True
            run.font.color.rgb = RGBColor.from_string(BLUE)


def grouped_notices(rows):
    groups = {}
    allowed = {'整体', '标题', '描述', '后台搜索词'} | {f'五点 {i}' for i in range(1, 101)}
    for row in rows:
        prefix, marker, message = row[0].partition(': ')
        if not marker or prefix not in allowed:
            prefix, message = '整体', row[0]
        fields = groups.setdefault(message, [])
        if prefix not in fields:
            fields.append(prefix)
    return [[message, '、'.join(fields)] for message, fields in groups.items()]


def report_document(payload):
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.27), Inches(11.69)
    section.top_margin = section.bottom_margin = Inches(.65)
    section.left_margin = section.right_margin = Inches(.7)
    for name in ['Normal', 'Title', 'Heading 1', 'Heading 2']:
        style = doc.styles[name]
        style.font.name = 'Arial'
        style.font.underline = False
        style.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), 'Noto Sans CJK SC')
        style.font.color.rgb = RGBColor.from_string(NAVY if name.startswith('Heading') else '202939')
    doc.styles['Title'].font.color.rgb = RGBColor(0, 0, 0)
    doc.styles['Title'].font.size = Pt(29)
    doc.styles['Heading 1'].font.size = Pt(16)
    doc.styles['Heading 2'].font.size = Pt(12)
    doc.styles['Normal'].font.size = Pt(10)
    doc.styles['Normal'].paragraph_format.space_after = Pt(6)
    doc.styles['Normal'].paragraph_format.line_spacing = 1.12
    eyebrow = doc.add_paragraph('AMAZON US')
    eyebrow.runs[0].font.color.rgb = RGBColor.from_string(MUTED)
    eyebrow.runs[0].font.size = Pt(10)
    title = doc.add_paragraph(payload['title'], 'Title')
    borders = OxmlElement('w:pBdr')
    for edge in ['top', 'bottom', 'left', 'right', 'between', 'bar']:
        border = OxmlElement('w:' + edge)
        border.set(qn('w:val'), 'nil')
        borders.append(border)
    title._p.get_or_add_pPr().append(borders)
    subtitle = doc.add_paragraph(payload['subtitle'].replace(' / ', '  '))
    subtitle.runs[0].font.color.rgb = RGBColor.from_string(MUTED)
    doc.add_heading('交付概览', level=1)
    comparison = next((s for s in payload['sections'] if s['title'] == '文案对比'), None)
    facts = next((s for s in payload['sections'] if s['title'].startswith('产品事实')), None)
    if comparison:
        count = sum(bool(row[4]) for row in comparison['rows'])
        metrics = doc.add_table(rows=1, cols=3)
        for cell, text in zip(metrics.rows[0].cells, [f"{len(comparison['rows'])} 个文案字段", f'{count} 个字段可复制', f"{len(facts['rows']) if facts else 0} 条产品事实"]):
            cell.text = text
            shade(cell, LIGHT)
        table_style(metrics, [2.29] * 3)
        doc.add_paragraph()
    for index, text in enumerate(payload['summary'].split('\n')):
        if index == 1:
            doc.add_heading('修改方向', level=2)
            text = text.removeprefix('修改方向：')
        elif index == 2:
            doc.add_heading('需要核实', level=2)
            text = text.removeprefix('需要核实：')
        doc.add_paragraph(text)
    doc.add_heading('阅读顺序', level=2)
    doc.add_paragraph('先查看文案对比，再核对产品事实与竞品参考。完整检查事项及版本记录见附录。')
    doc.core_properties.author = ''
    doc.core_properties.title = payload['title']
    footer = section.footer.paragraphs[0]
    footer.add_run('运营参考稿    美国站英文    ')
    page = OxmlElement('w:fldSimple')
    page.set(qn('w:instr'), 'PAGE')
    footer._p.append(page)
    for run in footer.runs:
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor.from_string(MUTED)
    for sec in payload['sections']:
        if sec['title'] in ['文案对比', '核实事项'] or (sec['title'] == '竞品参考' and sec['rows'] and sec['rows'][0][0] != '未取得竞品记录'):
            doc.add_page_break()
        doc.add_heading(sec['title'] if sec['title'] not in ['核实事项', '版本与规则'] else '附录 ' + sec['title'], level=1)
        if not sec['rows'] and sec['title'] != '核实事项':
            doc.add_paragraph('无记录')
            continue
        if sec['title'] == '文案对比':
            doc.add_paragraph('蓝色粗体标出当前参考稿相对原页面或上一版的新增与替换文字。文案内容不因视觉标注而改变。')
            for row in sec['rows']:
                start_paragraph = len(doc.paragraphs)
                field, original, previous, current, copyable, reason, sources, checks = row
                doc.add_heading(field, level=2)
                use_original = original not in ['新品，无原页面', '未取得原页面字段']
                before = original if use_original else previous
                table = doc.add_table(rows=2, cols=2)
                table.rows[0].cells[0].text = '原页面 未经独立核实' if use_original else '上一版草稿'
                table.rows[0].cells[1].text = '当前参考稿'
                table.rows[1].cells[0].text = before
                changed_text(table.rows[1].cells[1].paragraphs[0], before, current)
                for cell in table.rows[0].cells:
                    shade(cell, NAVY)
                    for run in cell.paragraphs[0].runs:
                        run.font.color.rgb = RGBColor.from_string('FFFFFF')
                        run.bold = True
                shade(table.rows[1].cells[0], LIGHT)
                shade(table.rows[1].cells[1], 'EEF4FB')
                table_style(table, [3.435, 3.435])
                doc.add_paragraph(('上一版草稿：' + previous) if use_original else ('原页面：' + original))
                p = doc.add_paragraph()
                p.add_run('修改原因  ').bold = True
                p.add_run(reason)
                p = doc.add_paragraph()
                p.add_run('产品依据  ').bold = True
                p.add_run(sources or '未取得依据')
                p = doc.add_paragraph('已进入 Excel 可复制列' if copyable else '可复制文案留空  当前仅供参考  检查事项见附录')
                p.runs[0].font.color.rgb = RGBColor.from_string(BLUE if copyable else '8A5A13')
                # Positive checks have no verification appendix entry, so preserve their exact wording here.
                if copyable:
                    doc.add_paragraph(checks)
                # Keep ordinary comparison cards together; unusually long material can still flow.
                if sum(len(value) for value in row) < 2500:
                    for paragraph in doc.paragraphs[start_paragraph:-1]:
                        paragraph.paragraph_format.keep_with_next = True
                    for table_row in table.rows:
                        for cell in table_row.cells:
                            for paragraph in cell.paragraphs:
                                paragraph.paragraph_format.keep_with_next = True
        elif sec['title'].startswith('产品事实'):
            keep_small_table(source_table(doc, sec['headers'], sec['rows'], [.75, 1.7, .85, 2.35, 1.22]), sec['rows'])
        elif sec['title'] == '核实事项':
            doc.add_paragraph('相同检查事项集中列示，涉及字段单独标注。待核实内容不会因报告排版改变而视为已解决。')
            check_rows = list(sec['rows'])
            if comparison:
                for row in comparison['rows']:
                    if not row[4]:
                        check_rows.extend([[row[0] + ': ' + message] for message in row[7].split('\n') if message])
            if check_rows:
                groups = grouped_notices(check_rows)
                display = [[message, '整体' if '整体' in fields.split('、') else fields] for message, fields in groups]
                source_table(doc, ['需核实内容', '涉及字段'], display, [5.2, 1.67])
            else:
                doc.add_paragraph('无待核实事项')
        elif sec['title'] == '版本与规则':
            keep_small_table(source_table(doc, sec['headers'], sec['rows'], [1.3, 5.57]), sec['rows'])
        elif sec['title'] == '竞品参考' and sec['rows'][0][0] == '未取得竞品记录':
            doc.add_paragraph(sec['rows'][0][0])
            doc.add_paragraph(sec['rows'][0][-1])
        else:
            # Competitor titles and sources remain full text, with no compressed eight-column table.
            for row in sec['rows']:
                doc.add_heading(row[0], level=2)
                for header, value in zip(sec['headers'][1:], row[1:]):
                    if value:
                        p = doc.add_paragraph()
                        p.add_run(header + '  ').bold = True
                        p.add_run(value)
    return doc


def render(payload, directory):
    directory = Path(directory)
    directory.mkdir(parents=True, exist_ok=True)
    stem = 'listing-' + hashlib.sha256((RENDER_VERSION + '\n' + json.dumps(payload, ensure_ascii=False, separators=(',', ':'))).encode()).hexdigest()
    doc = report_document(payload)
    wb = Workbook()
    wb.properties.creator = ''
    wb.properties.created = wb.properties.modified = datetime(2000, 1, 1)
    wb.remove(wb.active)
    for sec in payload['sections']:
        sheet = wb.create_sheet(sec['title'][:31])
        sheet.sheet_view.showGridLines = False
        sheet.append([sec['title']])
        sheet.append([payload['subtitle']])
        sheet.append(sec['headers'])
        for row in sec['rows']:
            sheet.append(row)
        sheet.freeze_panes = 'B4'
        sheet.auto_filter.ref = f'A3:{sheet.cell(3, len(sec["headers"])).column_letter}{max(sheet.max_row, 3)}'
        for row in sheet:
            for cell in row:
                if cell.value is not None:
                    cell.value = str(cell.value)
                    cell.data_type = 's'
                cell.number_format = '@'
                cell.font = Font(name='Arial', size=11, color='202020')
                cell.alignment = Alignment(wrap_text=True, vertical='top')
                if cell.row == 3:
                    cell.font = Font(name='Arial', size=11, bold=True, color='FFFFFF')
                    cell.fill = PatternFill('solid', fgColor='334155')
                elif cell.row == 1:
                    cell.font = Font(name='Arial', size=16, bold=True, color=NAVY)
                elif cell.row == 2:
                    cell.font = Font(name='Arial', size=10, color=MUTED)
                else:
                    if cell.row % 2 == 0:
                        cell.fill = PatternFill('solid', fgColor=LIGHT)
                    cell.border = Border(bottom=Side(style='hair', color='E4E9F0'))
                    if sec['title'] == '文案对比' and cell.column == 4:
                        cell.fill = PatternFill('solid', fgColor='EEF4FB')
                        cell.font = Font(name='Arial', size=11, color=NAVY, bold=True)
                    if sec['title'] == '文案对比' and cell.column == 5:
                        cell.fill = PatternFill('solid', fgColor='EDF5F0' if cell.value else AMBER)
        for column in sheet.columns:
            letter = column[0].column_letter
            sheet.column_dimensions[letter].width = 100 if len(sec['headers']) == 1 else 22 if letter == 'A' else 55
        if sec['title'] == '文案对比':
            for letter in ['B', 'C']:
                sheet.column_dimensions[letter].width = 40
            for letter in ['D', 'E']:
                sheet.column_dimensions[letter].width = 62
            # Supporting evidence is still visible; operators can collapse it when copying.
            for letter in ['F', 'G', 'H']:
                sheet.column_dimensions[letter].outlineLevel = 1
        if len(sec['headers']) > 1:
            sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(sec['headers']))
            sheet.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(sec['headers']))
        sheet.row_dimensions[1].height = 28
        sheet.row_dimensions[2].height = 35
        sheet.row_dimensions[3].height = 40
        for i, row in enumerate(sec['rows'], 4):
            # Long source excerpts remain readable by expanding rows, never truncating text.
            def line_count(text, col):
                width = sheet.column_dimensions[sheet.cell(3, col).column_letter].width * .85
                return sum(max(1, int((sum(2 if unicodedata.east_asian_width(c) in 'WF' else 1 for c in line) + width - 1) // width)) for line in text.split('\n'))
            lines = max((line_count(str(value), col) for col, value in enumerate(row, 1)), default=1)
            sheet.row_dimensions[i].height = min(409, max(40, lines * 16 + 8))
    docx, xlsx = directory / (stem + '.docx'), directory / (stem + '.xlsx')
    # Stable filename and exclusive output directory are controlled by the worker, not external text.
    # Canonical ZIP timestamps and immutable outputs prevent duplicate sends on rerender/retry.
    with tempfile.TemporaryDirectory() as temp:
        for target, save in [(docx, doc.save), (xlsx, wb.save)]:
            raw = Path(temp) / ('raw' + target.suffix)
            canonical = Path(temp) / ('canonical' + target.suffix)
            save(raw)
            with zipfile.ZipFile(raw) as source, zipfile.ZipFile(canonical, 'w', zipfile.ZIP_DEFLATED) as output:
                for name in sorted(source.namelist()):
                    info = zipfile.ZipInfo(name, (2000, 1, 1, 0, 0, 0))
                    info.compress_type = zipfile.ZIP_DEFLATED
                    content = source.read(name)
                    if name == 'docProps/core.xml':
                        content = re.sub(rb'(<dcterms:(?:created|modified)[^>]*>).*?(</dcterms:(?:created|modified)>)', rb'\g<1>2000-01-01T00:00:00Z\g<2>', content)
                    output.writestr(info, content)
            # Both locations must share a filesystem for hard links; stage beside final file.
            handle, stage = tempfile.mkstemp(prefix='.listing-', dir=directory)
            try:
                with os.fdopen(handle, 'wb') as stream:
                    stream.write(canonical.read_bytes())
                try:
                    os.link(stage, target)
                except FileExistsError:
                    if Path(stage).read_bytes() != target.read_bytes():
                        raise ValueError('LISTING_EXPORT_OUTPUT_CONFLICT')
            finally:
                Path(stage).unlink(missing_ok=True)
    return {'docx': str(docx), 'xlsx': str(xlsx)}


if __name__ == '__main__':
    payload = json.loads(Path(sys.argv[1]).read_text(encoding='utf-8'))
    print(json.dumps(render(payload, sys.argv[2])))
