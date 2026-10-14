"""Extract the actual deliverable, with its digest, before publication.

Runs with the task's OS identity. Never treats truncation as a full read.
"""
import hashlib
import io
import json
import os
import pathlib
import sys

# Third-party parsers can emit native warnings directly to fd 1. Reserve the
# original stdout for the protocol result; diagnostics belong on stderr.
protocol = os.fdopen(os.dup(sys.stdout.fileno()), 'w', encoding='utf-8')
sys.stdout.flush()
os.dup2(sys.stderr.fileno(), sys.stdout.fileno())

path = pathlib.Path(sys.argv[1])
raw = path.read_bytes()
ext = path.suffix.lower()
text = None
if ext == '.pdf':
    import fitz
    with fitz.open(stream=raw, filetype='pdf') as doc:
        text = '\n'.join(f'[Page {i+1}]\n{page.get_text(sort=True)}' for i, page in enumerate(doc))
elif ext == '.docx':
    from docx import Document
    doc = Document(io.BytesIO(raw))
    text = '\n'.join([p.text for p in doc.paragraphs] + [' | '.join(c.text for c in r.cells) for t in doc.tables for r in t.rows])
elif ext == '.xlsx':
    import openpyxl
    book = openpyxl.load_workbook(io.BytesIO(raw), data_only=True, read_only=True)
    text = '\n'.join(f'[{sheet.title}:{i}] ' + ' | '.join(str(c) if c is not None else '' for c in row)
                     for sheet in book for i, row in enumerate(sheet.values, 1) if any(c is not None for c in row))
    book.close()
elif ext == '.pptx':
    from pptx import Presentation
    doc = Presentation(io.BytesIO(raw))
    text = '\n'.join(f'[Slide {i+1}] ' + shape.text for i, slide in enumerate(doc.slides) for shape in slide.shapes if shape.has_text_frame)
elif ext in ('.txt', '.md', '.csv', '.json', '.html'):
    text = raw.decode('utf-8-sig')
if text is not None and len(text) > 160000:
    raise ValueError('OUTPUT_REVIEW_CONTENT_TOO_LARGE: split the deliverable into independently reviewable files')
print(json.dumps({'sha256': hashlib.sha256(raw).hexdigest(), 'text': text, 'complete': text is not None}, ensure_ascii=False), file=protocol, flush=True)
