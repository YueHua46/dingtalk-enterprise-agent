import importlib.util
import json
import base64
import io
import os
import tempfile
import unittest
from unittest.mock import patch
from PIL import Image
from openpyxl import Workbook
from openpyxl.drawing.image import Image as WorkbookImage
from openpyxl.styles import PatternFill

spec = importlib.util.spec_from_file_location("knowledge_parse", os.path.join(os.path.dirname(__file__), "knowledge_parse.py"))
parser = importlib.util.module_from_spec(spec)
spec.loader.exec_module(parser)


class KnowledgeParsingTests(unittest.TestCase):
    def test_svg_renders_full_extent_before_the_existing_ocr_pipeline(self):
        with tempfile.TemporaryDirectory() as directory:
            path = os.path.join(directory, "drawing.svg")
            with open(path, "w") as handle:
                handle.write('<svg xmlns="http://www.w3.org/2000/svg" width="60" height="40"><rect width="60" height="40" fill="red"/><rect x="50" y="30" width="10" height="10" fill="blue"/></svg>')
            image = parser.svg_image(path)
            self.assertEqual(image.size, (120, 80))
            self.assertEqual(image.getpixel((5, 5)), (255, 0, 0))
            self.assertEqual(image.getpixel((115, 75)), (0, 0, 255))
            manifest = os.path.join(directory, "manifest")
            os.mkdir(manifest)
            with patch("pytesseract.image_to_data", side_effect=AssertionError("manifest must not OCR")):
                result = parser.parse(path, manifest)
            self.assertEqual(len(result["images"]), 1)
            self.assertEqual(result["images"][0]["position"], {"image": "inline-svg"})

    def test_svg_rejects_active_external_and_unsupported_content(self):
        cases = [('KNOWLEDGE_SVG_UNSAFE', '<!DOCTYPE svg [<!ENTITY x SYSTEM "file:///etc/passwd">]><svg/>'),
                 ('KNOWLEDGE_SVG_UNSUPPORTED', '<svg><script/></svg>'),
                 ('KNOWLEDGE_SVG_UNSUPPORTED', '<svg><foreignObject/></svg>'),
                 ('KNOWLEDGE_SVG_UNSAFE', '<svg><use href="https://example.invalid/image"/></svg>'),
                 ('KNOWLEDGE_SVG_UNSAFE', '<svg><use href="file:///etc/passwd"/></svg>'),
                 ('KNOWLEDGE_SVG_UNSAFE', '<svg><rect style="fill:url(https://example.invalid/image)"/></svg>'),
                 ('KNOWLEDGE_SVG_UNSAFE', '<svg><style>@import "https://example.invalid/css";</style></svg>'),
                 ('KNOWLEDGE_SVG_UNSAFE', '<svg><style>rect{fill:u\\72l(file:///secret)}</style></svg>'),
                 ('KNOWLEDGE_SVG_UNSUPPORTED', '<svg xmlns="https://example.invalid/namespace"/>'),
                 ('KNOWLEDGE_SVG_UNSAFE', '<svg onload="run()"/>')]
        with tempfile.TemporaryDirectory() as directory:
            path = os.path.join(directory, "drawing.svg")
            for code, content in cases:
                with self.subTest(code=code, content=content):
                    with open(path, "w") as handle:
                        handle.write(content)
                    with self.assertRaisesRegex(ValueError, code):
                        parser.svg_image(path)
            with open(path, "wb") as handle:
                handle.write('<!DOCTYPE svg [<!ENTITY x SYSTEM "file:///etc/passwd">]><svg/>'.encode("utf-16"))
            with self.assertRaisesRegex(ValueError, "KNOWLEDGE_SVG_UNSAFE"):
                parser.svg_image(path)

    def test_svg_rejects_excessive_raster_extent(self):
        with tempfile.TemporaryDirectory() as directory:
            path = os.path.join(directory, "large.svg")
            with open(path, "w") as handle:
                handle.write('<svg xmlns="http://www.w3.org/2000/svg" width="100000" height="100000"><rect width="1" height="1"/></svg>')
            with self.assertRaisesRegex(ValueError, "KNOWLEDGE_SVG_SIZE_LIMIT"):
                parser.svg_image(path)

    def test_regions_cover_bottom_and_map_coordinates_without_merging_lines(self):
        seen = []
        def ocr(image, **options):
            seen.append(image.size)
            return {"text": ["password: fixture"], "block_num": [1], "par_num": [1], "line_num": [1], "top": [20], "height": [12]}
        with patch("pytesseract.image_to_data", side_effect=ocr):
            words = parser.credential_ocr(Image.new("RGB", (40, 1100)), 2)
        self.assertEqual(seen, [(80, 800), (80, 800), (80, 800), (80, 184)])
        self.assertEqual(words["top"], [20, 692, 1364, 2036])
        self.assertEqual(len(set(words["block_num"])), 4)

    def test_later_region_credentials_are_masked_at_original_location(self):
        empty = {key: [] for key in ("text", "block_num", "par_num", "line_num", "top", "height")}
        later = {"text": ["password: fixture"], "block_num": [1], "par_num": [1], "line_num": [1], "top": [28], "height": [12]}
        with patch("pytesseract.image_to_data", side_effect=[empty, later]):
            result = parser.encoded(Image.new("RGB", (40, 700), "black"), {}, True)
        image = Image.open(io.BytesIO(base64.b64decode(result["dataUrl"].split(",")[1])))
        self.assertEqual(result["redactions"], 1)
        self.assertGreater(image.getpixel((20, 350))[0], 240)
        self.assertLess(image.getpixel((20, 100))[0], 10)

    def test_regions_share_one_deadline_and_reject_late_results(self):
        empty = {key: [] for key in ("text", "block_num", "par_num", "line_num", "top", "height")}
        with patch("pytesseract.image_to_data", return_value=empty) as ocr, patch.object(parser.time, "monotonic", side_effect=[0, 1, 50, 51, 91]):
            with self.assertRaisesRegex(ValueError, "KNOWLEDGE_OCR_TIMEOUT"):
                parser.encoded(Image.new("RGB", (40, 700)), {}, True)
        self.assertEqual([round(call.kwargs["timeout"]) for call in ocr.call_args_list], [89, 39])

    def test_one_failed_region_rejects_complete_image(self):
        empty = {key: [] for key in ("text", "block_num", "par_num", "line_num", "top", "height")}
        with patch("pytesseract.image_to_data", side_effect=[empty, RuntimeError("OCR failed")]):
            with self.assertRaisesRegex(ValueError, "KNOWLEDGE_OCR_FAILED"):
                parser.encoded(Image.new("RGB", (40, 700)), {}, True)

    def test_sparse_formatting_does_not_hide_late_data_or_expand_rectangle(self):
        with tempfile.TemporaryDirectory() as directory:
            book = Workbook()
            sheet = book.active
            sheet["A1"] = "start"
            sheet["C500000"] = "late value"
            sheet["XFD1048576"].fill = PatternFill("solid", fgColor="FFFFFF")
            path = os.path.join(directory, "sparse.xlsx")
            book.save(path)
            manifest = os.path.join(directory, "manifest")
            os.mkdir(manifest)
            result = parser.parse(path, manifest)
            self.assertEqual(result["sheets"][0]["rows"], [{"row": 1, "cells": ["start"]}, {"row": 500000, "cells": ["", "", "late value"]}])

    def test_all_105_images_are_manifested_and_paged_without_ocr_in_manifest(self):
        with tempfile.TemporaryDirectory() as directory:
            image_path = os.path.join(directory, "image.png")
            Image.new("RGB", (16, 16), "white").save(image_path)
            book = Workbook()
            for index in range(105):
                book.active.add_image(WorkbookImage(image_path), "C" + str(index + 1))
            path = os.path.join(directory, "images.xlsx")
            book.save(path)
            manifest = os.path.join(directory, "manifest")
            os.mkdir(manifest)
            with patch("pytesseract.image_to_data", side_effect=AssertionError("manifest must not OCR")):
                result = parser.parse(path, manifest)
            self.assertEqual(len(result["images"]), 105)
            self.assertEqual(result["images"][-1]["position"]["range"], "C105")
            self.assertTrue(all("dataUrl" not in entry for entry in result["images"]))
            with open(os.path.join(manifest, "manifest.json"), "w") as handle:
                json.dump(result, handle)
            empty = {key: [] for key in ("text", "block_num", "par_num", "line_num", "top", "height")}
            rendered = []
            with patch("pytesseract.image_to_data", return_value=empty):
                for cursor in range(0, 105, 4):
                    batch = parser.image_batch(manifest, cursor, 4)
                    self.assertEqual(batch["total"], 105)
                    self.assertEqual(batch["sourceHash"], result["sourceHash"])
                    rendered.extend(batch["images"])
            self.assertEqual([i["hash"] for i in rendered], [i["hash"] for i in result["images"]])
            self.assertEqual(len(rendered), 105)
            self.assertTrue(all(i["dataUrl"].startswith("data:image/jpeg;base64,") for i in rendered))
            with self.assertRaisesRegex(ValueError, "KNOWLEDGE_IMAGE_BATCH_INVALID"):
                parser.image_batch(manifest, 0, 5)

    def test_ocr_timeout_fails_closed_without_unredacted_upload(self):
        with patch("pytesseract.image_to_data", side_effect=RuntimeError("Tesseract process timeout")):
            with self.assertRaisesRegex(ValueError, "KNOWLEDGE_OCR_TIMEOUT"):
                parser.encoded(Image.new("RGB", (20, 20)), {}, True)

    def test_normal_password_steps_are_kept_but_credential_values_are_masked(self):
        def words(text):
            return {"text": [text], "block_num": [1], "par_num": [1], "line_num": [1], "top": [1], "height": [8]}
        with patch("pytesseract.image_to_data", return_value=words("点击修改密码并输入新密码")):
            self.assertEqual(parser.encoded(Image.new("RGB", (40, 40)), {}, True)["redactions"], 0)
        with patch("pytesseract.image_to_data", return_value=words("密码：fixture-value")):
            self.assertEqual(parser.encoded(Image.new("RGB", (40, 40)), {}, True)["redactions"], 1)


if __name__ == "__main__":
    unittest.main()
