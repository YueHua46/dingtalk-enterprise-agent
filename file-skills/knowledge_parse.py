"""Bounded document decoding in the existing per-task worker sandbox.

No provider credentials, network requests or macros. Images returned to the
host are explicit source evidence, with page/cell locations.
"""
import base64
import io
import json
import os
import sys
import zipfile
import hashlib
import time
from PIL import Image
from PIL import ImageDraw
import re
from concurrent.futures import ThreadPoolExecutor

Image.MAX_IMAGE_PIXELS = 40_000_000
MAX_CELLS = 2_000_000
MAX_IMAGES = 100
RENDER_IMAGES = len(sys.argv) < 3 or sys.argv[2] == "true"
os.environ["OMP_THREAD_LIMIT"] = "1"


def credential_ocr(image, scale):
    """Full-width overlapping regions bound dense layout analysis, not content.

    All regions must succeed inside one image's original 90-second budget.
    Return coordinates in scaled full-image space for existing line masking.
    """
    import pytesseract
    deadline = time.monotonic() + 90
    combined = {key: [] for key in ("text", "block_num", "par_num", "line_num", "top", "height")}
    top = 0
    while top < image.height:
        bottom = min(top + 400, image.height)
        region = image.crop((0, top, image.width, bottom))
        if scale > 1:
            region = region.resize((region.width * scale, region.height * scale))
        remaining = deadline - time.monotonic()
        if remaining <= 0:
            raise ValueError("KNOWLEDGE_OCR_TIMEOUT")
        try:
            words = pytesseract.image_to_data(region, lang="eng+chi_sim", config="--psm 11", output_type=pytesseract.Output.DICT, timeout=remaining)
        except RuntimeError as error:
            raise ValueError("KNOWLEDGE_OCR_TIMEOUT" if "timeout" in str(error).lower() else "KNOWLEDGE_OCR_FAILED") from None
        if time.monotonic() > deadline:
            raise ValueError("KNOWLEDGE_OCR_TIMEOUT")
        for key in combined:
            if key == "top":
                combined[key].extend(value + top * scale for value in words[key])
            elif key == "block_num":
                combined[key].extend((top, value) for value in words[key])
            else:
                combined[key].extend(words[key])
        if bottom == image.height:
            break
        top = bottom - 64
    return combined


def encoded(image, position, render=None):
    image = image.convert("RGB")
    source_hash = hashlib.sha256(image.tobytes() + str(image.size).encode()).hexdigest()
    if not (RENDER_IMAGES if render is None else render):
        return {"hash": source_hash, "position": position}
    # Local OCR identifies credential-bearing lines before provider upload.
    # Erase the entire matching line (labels and values), not just the label.
    scale = 2 if image.width * image.height * 4 <= Image.MAX_IMAGE_PIXELS else 1
    words = credential_ocr(image, scale)
    lines = {}
    for i, text in enumerate(words["text"]):
        key = (words["block_num"][i], words["par_num"][i], words["line_num"][i])
        lines.setdefault(key, []).append(i)
    draw = ImageDraw.Draw(image)
    redactions = 0
    for indices in lines.values():
        text = " ".join(words["text"][i] for i in indices)
        compact = re.sub(r"[^a-z0-9\u4e00-\u9fff]", "", text.lower())
        password_instruction = re.search(r"(?:修改|重置|输入|设置|更改)(?:新|原|登录)?密码|(?:change|reset|enter|set)(?:your|the|new)?password", compact)
        password_value = re.search(r"(?:password|passwd|密码|口令)\s*[:：=]\s*\S+", text, re.I)
        credential_label = re.search(r"apikey|appsecret|clientsecret|accesstoken|refreshtoken|stream(?:key|id|ld|1d)|wssecret|密钥|验证码", compact)
        credential_password = re.search(r"password|passwd|密码|口令", compact) and (password_value or not password_instruction)
        if credential_label or credential_password or re.search(r"\b[a-z0-9]{4}(?:\s*-\s*[a-z0-9]{4}){3,}\b|[a-f0-9]{24,}|\bsk-[a-z0-9_-]{8,}", text, re.I):
            top = min(words["top"][i] for i in indices) // scale
            bottom = max(words["top"][i] + words["height"][i] for i in indices) // scale
            draw.rectangle((0, max(0, top - 6), image.width, min(image.height, bottom + 30)), fill="white")
            redactions += 1
    image.thumbnail((1800, 1800))
    out = io.BytesIO()
    image.save(out, format="JPEG", quality=85)
    return {"hash": source_hash, "redactions": redactions, "dataUrl": "data:image/jpeg;base64," + base64.b64encode(out.getvalue()).decode(), "position": position}


def parse(path, manifest_dir=None):
    ext = os.path.splitext(path)[1].lower()
    result = {"sheets": [], "pages": [], "images": [], "quality": []}
    def add_image(image, position):
        if manifest_dir:
            entry = encoded(image, position, False)
            entry["file"] = str(len(result["images"])) + ".png"
            image.convert("RGB").save(os.path.join(manifest_dir, entry["file"]), format="PNG")
            result["images"].append(entry)
        else:
            result["images"].append(encoded(image, position))
    if ext == ".xlsx":
        import openpyxl
        # Bound expanded ZIP size before allocating workbook/media objects.
        with zipfile.ZipFile(path) as archive:
            if sum(e.file_size for e in archive.infolist()) > 250 * 1024 * 1024:
                raise ValueError("KNOWLEDGE_WORKBOOK_EXPANDED_LIMIT")
        book = openpyxl.load_workbook(path, data_only=True)
        for sheet in book.worksheets:
            # Worksheet dimensions include styled but empty remote cells.
            # Traverse stored values instead of allocating that rectangle.
            values = sorted((c for c in sheet._cells.values() if c.value is not None), key=lambda c: (c.row, c.column))
            if len(values) > MAX_CELLS:
                raise ValueError("KNOWLEDGE_WORKBOOK_CELL_LIMIT")
            rows = []
            by_row = {}
            for cell in values:
                by_row.setdefault(cell.row, {})[cell.column] = str(cell.value)
            if sum(max(cells) for cells in by_row.values()) > MAX_CELLS:
                raise ValueError("KNOWLEDGE_WORKBOOK_CELL_LIMIT")
            for index, cells in by_row.items():
                rows.append({"row": index, "cells": [cells.get(col, "") for col in range(1, max(cells) + 1)]})
            result["sheets"].append({"title": sheet.title, "rows": rows,
                                     "mergedRanges": [str(r) for r in sheet.merged_cells.ranges]})
            for index, media in enumerate(sheet._images):
                anchor = getattr(media.anchor, "_from", None)
                position = {"sheet": sheet.title, "image": str(index)}
                if anchor:
                    from openpyxl.utils import get_column_letter
                    position["range"] = get_column_letter(anchor.col + 1) + str(anchor.row + 1)
                add_image(Image.open(io.BytesIO(media._data())), position)
        book.close()
    elif ext == ".pdf":
        import pymupdf as fitz
        doc = fitz.open(path)
        if len(doc) > 1000:
            raise ValueError("KNOWLEDGE_PDF_PAGE_LIMIT")
        for index, page in enumerate(doc):
            blocks = page.get_text("blocks", sort=True)
            text = "\n".join(b[4] for b in blocks if len(b) > 6 and b[6] == 0)
            bad = len(text.strip()) < 50 or text.count("\ufffd") > max(2, len(text) // 100)
            result["pages"].append({"page": index + 1, "text": text,
                                    "regions": [{"box": list(b[:4]), "text": b[4]} for b in blocks if len(b) > 6 and b[6] == 0]})
            if bad or page.get_images() or page.get_drawings() or page.rect.width > page.rect.height * 1.2:
                # Crop wide multi-panel pages so labels remain legible.
                panels = 3 if page.rect.width > page.rect.height * 1.5 else 1
                for panel in range(panels):
                    box = fitz.Rect(page.rect.width * panel / panels, 0, page.rect.width * (panel + 1) / panels, page.rect.height)
                    pix = page.get_pixmap(matrix=fitz.Matrix(1.5, 1.5), clip=box)
                    add_image(Image.open(io.BytesIO(pix.tobytes("png"))), {"page": index + 1, "image": str(panel), "region": list(box)})
                if bad:
                    result["quality"].append("low_text_page:" + str(index + 1))
        doc.close()
    elif ext == ".svg":
        add_image(svg_image(path), {"image": "inline-svg"})
    elif ext in (".png", ".jpg", ".jpeg", ".bmp"):
        add_image(Image.open(path), {"image": "original"})
    else:
        raise ValueError("KNOWLEDGE_BINARY_UNSUPPORTED")
    if not manifest_dir and RENDER_IMAGES and len(result["images"]) > MAX_IMAGES:
        raise ValueError("KNOWLEDGE_IMAGE_LIMIT")
    result["sourceHash"] = hashlib.sha256(json.dumps({**result, "images": [{"hash": i["hash"], "position": i["position"]} for i in result["images"]]}, sort_keys=True, ensure_ascii=False).encode()).hexdigest()
    return result


def svg_image(path):
    """Decode a complete, self-contained imported drawing, without fetching.

    Reject active/external/unsupported content instead of omitting it. The
    existing MuPDF renderer handles vector and text geometry in the sandbox.
    """
    import math
    import xml.etree.ElementTree as ET
    import pymupdf as fitz
    with open(path, "rb") as handle:
        data = handle.read(8 * 1024 * 1024 + 1)
    if len(data) > 8 * 1024 * 1024:
        raise ValueError("KNOWLEDGE_SVG_UNSAFE")
    try:
        text = data.decode("utf-8-sig")
    except UnicodeDecodeError:
        raise ValueError("KNOWLEDGE_SVG_UNSAFE") from None
    if re.search(r"<!DOCTYPE|<!ENTITY", text, re.I):
        raise ValueError("KNOWLEDGE_SVG_UNSAFE")
    try:
        root = ET.fromstring(text)
    except ET.ParseError:
        raise ValueError("KNOWLEDGE_SVG_INVALID") from None
    allowed = {"svg", "g", "defs", "path", "rect", "circle", "ellipse", "line", "polyline", "polygon", "text", "tspan", "use", "clipPath", "linearGradient", "radialGradient", "stop", "title", "desc", "style"}
    for element in root.iter():
        if element.tag.split("}")[-1] not in allowed or (element.tag.startswith("{") and not element.tag.startswith("{http://www.w3.org/2000/svg}")):
            raise ValueError("KNOWLEDGE_SVG_UNSUPPORTED")
        for key, value in element.attrib.items():
            name = key.split("}")[-1].lower()
            if name.startswith("on") or name == "base" or (name == "href" and not re.fullmatch(r"#[A-Za-z_][A-Za-z0-9_.:-]*", value)):
                raise ValueError("KNOWLEDGE_SVG_UNSAFE")
        for value in [*element.attrib.values(), element.text or ""]:
            if re.search(r"@import", value, re.I) or any(not re.fullmatch(r"\s*['\"]?#[A-Za-z_][A-Za-z0-9_.:-]*['\"]?\s*", target) for target in re.findall(r"url\s*\((.*?)\)", value, re.I | re.S)):
                raise ValueError("KNOWLEDGE_SVG_UNSAFE")
        for value in [element.attrib.get("style", ""), element.text or "" if element.tag.split("}")[-1] == "style" else ""]:
            if "\\" in value or "@" in value:
                raise ValueError("KNOWLEDGE_SVG_UNSAFE")
    if root.tag.split("}")[-1] != "svg":
        raise ValueError("KNOWLEDGE_SVG_INVALID")
    with fitz.open(stream=data, filetype="svg") as doc:
        if len(doc) != 1:
            raise ValueError("KNOWLEDGE_SVG_INVALID")
        page = doc[0]
        width, height = page.rect.width, page.rect.height
        if not all(math.isfinite(v) and v > 0 for v in (width, height)) or math.ceil(width * 2) * math.ceil(height * 2) > Image.MAX_IMAGE_PIXELS:
            raise ValueError("KNOWLEDGE_SVG_SIZE_LIMIT")
        pix = page.get_pixmap(matrix=fitz.Matrix(2, 2), alpha=False)
        return Image.open(io.BytesIO(pix.tobytes("png")))


def image_batch(directory, cursor, count):
    if cursor < 0 or not 1 <= count <= 4:
        raise ValueError("KNOWLEDGE_IMAGE_BATCH_INVALID")
    with open(os.path.join(directory, "manifest.json"), encoding="utf-8") as handle:
        manifest = json.load(handle)
    entries = manifest["images"]
    if cursor > len(entries):
        raise ValueError("KNOWLEDGE_IMAGE_BATCH_INVALID")
    def render(entry):
        if not re.fullmatch(r"\d+\.png", entry["file"]):
            raise ValueError("KNOWLEDGE_IMAGE_PATH_INVALID")
        with Image.open(os.path.join(directory, entry["file"])) as image:
            result = encoded(image, entry["position"], True)
        if result["hash"] != entry["hash"]:
            raise ValueError("KNOWLEDGE_IMAGE_CHANGED")
        return result
    with ThreadPoolExecutor(max_workers=4) as executor:
        images = list(executor.map(render, entries[cursor:cursor + count]))
    return {"images": images, "next": cursor + len(images), "total": len(entries), "sourceHash": manifest["sourceHash"]}


if __name__ == "__main__":
    try:
        if sys.argv[2:3] == ["manifest"]:
            directory = sys.argv[3]
            os.makedirs(directory, exist_ok=True)
            result = parse(sys.argv[1], directory)
            with open(os.path.join(directory, "manifest.json"), "w", encoding="utf-8") as handle:
                json.dump(result, handle, ensure_ascii=False)
            print(json.dumps(result, ensure_ascii=False))
        elif sys.argv[2:3] == ["batch"]:
            print(json.dumps(image_batch(sys.argv[1], int(sys.argv[3]), int(sys.argv[4])), ensure_ascii=False))
        else:
            print(json.dumps(parse(sys.argv[1]), ensure_ascii=False))
    except Exception as error:
        # Only stable classifications leave the sandbox; no raw text/paths.
        code = str(error) if re.fullmatch(r"KNOWLEDGE_[A-Z_0-9]+", str(error)) else "KNOWLEDGE_BINARY_INVALID"
        print(code, file=sys.stderr)
        sys.exit(1)
