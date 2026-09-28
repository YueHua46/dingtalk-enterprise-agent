"""Verify an exact excerpt against a task input document."""

import json
import re
import sys
from pathlib import Path


def normalize(value):
    return re.sub(r"\s+", " ", str(value)).strip().casefold()


def document_text(path, pages):
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        import pymupdf
        import pytesseract
        from PIL import Image

        document = pymupdf.open(path)
        selected = pages or list(range(1, min(len(document), 20) + 1))
        chunks = []
        for number in selected:
            if number < 1 or number > len(document):
                raise ValueError("PAGE_OUT_OF_RANGE")
            page = document[number - 1]
            text = page.get_text()
            if len(text.strip()) < 12:
                pixmap = page.get_pixmap(matrix=pymupdf.Matrix(2, 2))
                image = Image.frombytes("RGB", [pixmap.width, pixmap.height], pixmap.samples)
                text = pytesseract.image_to_string(image, lang="eng+chi_sim")
            chunks.append(text)
        return "\n".join(chunks), len(document)
    if suffix == ".docx":
        from docx import Document

        document = Document(path)
        chunks = [item.text for item in document.paragraphs]
        chunks.extend(cell.text for table in document.tables for row in table.rows for cell in row.cells)
        return "\n".join(chunks), None
    if suffix == ".xlsx":
        from openpyxl import load_workbook

        workbook = load_workbook(path, read_only=True, data_only=False)
        chunks = []
        for sheet in workbook:
            chunks.append(sheet.title)
            for row in sheet.iter_rows(max_row=min(sheet.max_row, 10000), max_col=min(sheet.max_column, 100)):
                chunks.append(" ".join(str(cell.value) for cell in row if cell.value is not None))
                chunks.extend(f"{cell.coordinate} {cell.value}" for cell in row if cell.value is not None)
        return "\n".join(chunks), None
    if suffix == ".pptx":
        from pptx import Presentation

        presentation = Presentation(path)
        chunks = []
        for slide in presentation.slides:
            for shape in slide.shapes:
                if shape.has_text_frame:
                    chunks.append(shape.text)
                if shape.has_chart:
                    chunks.extend(f"{series.name} {' '.join(map(str, series.values))}" for series in shape.chart.series)
        return "\n".join(chunks), len(presentation.slides)
    if suffix in {".txt", ".md", ".csv", ".json"}:
        return path.read_text(encoding="utf-8", errors="replace"), None
    raise ValueError("INPUT_TYPE_UNSUPPORTED")


path = Path(sys.argv[1])
excerpt = sys.argv[2]
pages = json.loads(sys.argv[3])
text, page_count = document_text(path, pages)
needle = normalize(excerpt)
print(json.dumps({"verified": bool(needle) and needle in normalize(text), "pageCount": page_count}))
