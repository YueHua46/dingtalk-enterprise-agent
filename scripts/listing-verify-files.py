"""Verify files downloaded from the original DingTalk conversation against host payload."""
import hashlib
import json
import sys
from pathlib import Path
from docx import Document
from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parent.parent / 'file-skills'))
from listing_export import grouped_notices

payload_path, docx_path, xlsx_path = map(Path, sys.argv[1:4])
payload = json.loads(payload_path.read_text(encoding='utf-8-sig'))
book = load_workbook(xlsx_path, data_only=False)
doc = Document(docx_path)
text = '\n'.join([p.text for p in doc.paragraphs] + [c.text for t in doc.tables for r in t.rows for c in r.cells])
cells = 0
for section in payload['sections']:
    sheet = book[section['title'][:31]]
    for r, row in enumerate(section['rows'], 4):
        for c, value in enumerate(row, 1):
            assert (sheet.cell(r, c).value or '') == value, (section['title'], r, c)
            assert sheet.cell(r, c).data_type != 'f', 'Formula detected'
            cells += 1
            if value:
                if section['title'] == '核实事项':
                    value = grouped_notices([[value]])[0][0]
                for line in value.split('\n'):
                    assert line in text, ('Word value missing', section['title'], line)
assert payload['versionId'] in text
assert len(book.sheetnames) == len(payload['sections'])
result = {'downloadedFilesMatchHostPayload': True, 'versionId': payload['versionId'],
          'checkedExcelCells': cells, 'sheets': len(book.sheetnames), 'formulaCells': 0,
          'files': {p.name: hashlib.sha256(p.read_bytes()).hexdigest() for p in [docx_path, xlsx_path]}}
print(json.dumps(result, ensure_ascii=False, indent=2))
