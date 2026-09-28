"""Run inside an isolated file-worker task directory to verify document runtimes."""

from pathlib import Path
import subprocess
import sys

from docx import Document
from openpyxl import Workbook, load_workbook
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches
from pypdf import PdfReader
from reportlab.pdfgen import canvas


root = Path(sys.argv[1]).resolve()
output = root / "output"
preview = root / "workspace" / "preview"
output.mkdir(exist_ok=True)
preview.mkdir(parents=True, exist_ok=True)

image = Image.new("RGB", (500, 160), "white")
ImageDraw.Draw(image).text((20, 55), "Image in document", fill="black")
image_path = preview / "image.png"
image.save(image_path)

doc = Document()
doc.add_heading("Document check", 0)
table = doc.add_table(rows=2, cols=2)
table.cell(0, 0).text = "Metric"
table.cell(0, 1).text = "Value"
table.cell(1, 0).text = "Orders"
table.cell(1, 1).text = "12"
doc.add_picture(str(image_path), width=Inches(3))
doc_path = output / "document.docx"
doc.save(doc_path)
reopened = Document(doc_path)
assert len(reopened.tables) == 1 and len(reopened.inline_shapes) == 1

book = Workbook()
sheet = book.active
sheet.title = "Sales"
sheet["A1"], sheet["B1"] = "Orders", 12
sheet["A2"], sheet["B2"] = "Price", 4
sheet["B3"] = "=B1*B2"
book.create_sheet("Notes")["A1"] = "Second sheet"
xlsx_path = output / "workbook.xlsx"
book.save(xlsx_path)
reopened_book = load_workbook(xlsx_path, data_only=False)
assert reopened_book.sheetnames == ["Sales", "Notes"]
assert reopened_book["Sales"]["B3"].value == "=B1*B2"

deck = Presentation()
slide = deck.slides.add_slide(deck.slide_layouts[5])
slide.shapes.title.text = "Quarterly orders"
chart_data = CategoryChartData()
chart_data.categories = ["Q1", "Q2"]
chart_data.add_series("Orders", [12, 18])
slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(1.5), Inches(7), Inches(4), chart_data)
pptx_path = output / "slides.pptx"
deck.save(pptx_path)
reopened_deck = Presentation(pptx_path)
assert len(reopened_deck.slides) == 1
assert any(shape.has_chart for shape in reopened_deck.slides[0].shapes)

pdf_path = output / "text.pdf"
pdf = canvas.Canvas(str(pdf_path))
pdf.drawString(72, 720, "Text PDF check")
pdf.save()
assert "Text PDF check" in PdfReader(pdf_path).pages[0].extract_text()

scan_image = Image.new("RGB", (1200, 300), "white")
font = ImageFont.truetype("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf", 60)
ImageDraw.Draw(scan_image).text((70, 100), "SCANNED PDF CHECK", fill="black", font=font)
scan_path = output / "scan.pdf"
scan_image.save(scan_path, "PDF", resolution=150.0)
assert not (PdfReader(scan_path).pages[0].extract_text() or "").strip()

for source in (doc_path, xlsx_path, pptx_path):
    subprocess.run([
        "libreoffice", f"-env:UserInstallation=file://{root}/tmp/lo-{source.stem}",
        "--headless", "--convert-to", "pdf", "--outdir", str(preview), str(source),
    ], check=True, timeout=90, capture_output=True)
    rendered = preview / (source.stem + ".pdf")
    assert rendered.exists() and len(PdfReader(rendered).pages) >= 1
    subprocess.run(["pdftoppm", "-f", "1", "-l", "1", "-png", "-r", "90", str(rendered), str(preview / source.stem)], check=True, timeout=90, capture_output=True)
    assert list(preview.glob(source.stem + "-*.png"))

subprocess.run(["pdftoppm", "-f", "1", "-l", "1", "-png", "-r", "150", str(scan_path), str(preview / "scan")], check=True, timeout=90, capture_output=True)
ocr = subprocess.run(["tesseract", str(next(preview.glob("scan-*.png"))), "stdout", "-l", "eng"], check=True, timeout=90, capture_output=True, text=True)
assert ocr.stdout.strip()
print("DOCX/XLSX/PPTX/text PDF/scanned PDF create, reopen, render and OCR: OK")
